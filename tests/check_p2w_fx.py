# PDF→Word regression checks on real-world fixtures (tester round 1). Usage: python3 check_p2w_fx.py DOCXDIR
import sys,os,re,json,zipfile,unicodedata
import docx
D=sys.argv[1];T=os.path.join(os.path.dirname(__file__),'fixtures','p2w_truth');ST=json.load(open(os.path.join(D,'_status.json')))
W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
def lev(a,b):
  try:
    from rapidfuzz.distance import Levenshtein;return Levenshtein.distance(a,b)
  except ImportError:
    import difflib;return round((1-difflib.SequenceMatcher(None,a,b,autojunk=False).ratio())*max(len(a),len(b)))
def norm(t):return re.sub(r'\s+',' ',unicodedata.normalize('NFC',t)).strip()
def text(p):
  d=docx.Document(p);return '\n'.join([x.text for x in d.paragraphs]+[' '.join(c.text for c in r.cells) for t in d.tables for r in t.rows])
def cer(p,truth):a=norm(text(p));b=norm(text(truth) if truth.endswith('.docx') else open(truth,encoding='utf-8').read());return lev(a,b)/len(b)
def xml(p):return zipfile.ZipFile(p).read('word/document.xml').decode()
fails=[]
def ok(c,msg):
  print(('PASS ' if c else 'FAIL ')+msg)
  if not c:fails.append(msg)
for pre in ['','m_']:
  f=lambda n:os.path.join(D,pre+n+'.docx')
  for n in ['vi_lo','vi_dejavu','Hợp đồng #1 (bản cuối) & thử']:
    c=cer(f(n),os.path.join(T,'vi_lo.docx'));ok(c<0.02,'%s%s CER %.4f'%(pre,n,c))
    x=xml(f(n));ok(x.count('<w:numPr>')>=6,'%s%s numPr %d'%(pre,n,x.count('<w:numPr>')))
    ok(x.count('<w:tbl>')==1,'%s%s only the real price table'%(pre,n))
    z=zipfile.ZipFile(f(n));ok('word/numbering.xml' in z.namelist() and 'docProps/core.xml' in z.namelist() and not [m for m in z.namelist() if m.endswith('/')],'%s%s package parts'%(pre,n))
    ok(len(re.findall(r'pageBreakBefore',x))==1,'%s%s 2 pages'%(pre,n))
  c=cer(f('book'),os.path.join(T,'book.txt'));ok(c<0.01,'%sbook CER %.4f'%(pre,c))
  x=xml(f('book'));ok(x.count('<w:tbl>')==0,'%sbook TOC is not a table'%pre);ok(x.count('<w:numPr>')==0,'%sbook no fake lists'%pre);ok(x.count('pageBreakBefore')==52,'%sbook 53 pages'%pre)
  x=xml(f('rot_a3'));sz=re.findall(r'<w:pgSz w:w="(\d+)" w:h="(\d+)"( w:orient="landscape")?',x)
  ok([(int(a)//20,int(b)//20,bool(o)) for a,b,o in sz]==[(842,595,True),(1191,842,True)],'%srot_a3 sections %s'%(pre,sz))
  x=xml(f('two_col'));t=norm(text(f('two_col')));ok(x.count('<w:tbl>')==0 and len(t)>2000,'%stwo_col text'%pre)
  ok(ST[pre+'two_col.pdf']['notes'] and 'multiple columns' in ST[pre+'two_col.pdf']['notes'][0],'%stwo_col multi-column note'%pre)
  s=ST[pre+'no_tounicode.pdf'];ok('garbled' in s['status'] and s['ocrhint'],'%sno_tounicode garbled warning + OCR link'%pre)
  s=ST[pre+'my_scan.pdf'];ok('scan' in s['status'] and s['ocrhint'],'%smy_scan scan hint'%pre)
  z=zipfile.ZipFile(f('my_scan'));ok(sorted(m for m in z.namelist() if 'media' in m)==['word/media/image1.jpeg','word/media/image2.jpeg'],'%smy_scan JPEG kept'%pre)
  for e in re.findall(r'<wp:extent cx="(\d+)" cy="(\d+)"',xml(f('my_scan'))):ok(int(e[1])/12700<=792-108,'%smy_scan image fits page height'%pre)
  c=cer(f('vi_list'),os.path.join(T,'vi_list.docx'));ok(c<0.005,'%svi_list CER %.4f'%(pre,c))
  ps=docx.Document(f('vi_list')).paragraphs;isl=lambda q:q._p.pPr is not None and q._p.pPr.numPr is not None
  ok(sum(map(isl,ps))==7,'%svi_list 7 real list items'%pre)
  for mk in ['1. Đây là đoạn văn thường','- Đây là đoạn văn thường']:
    hit=[q for q in ps if q.text.startswith(mk)];ok(len(hit)==1 and not isl(hit[0]) and 'trang' in hit[0].text[-30:],'%svi_list standalone "%s" kept as one plain paragraph'%(pre,mk[:2]))
  ok(xml(f('vi_list')).count('<w:tbl>')==1,'%svi_list real table'%pre)
  lv=[(q._p.pPr.numPr.ilvl.val,q._p.pPr.numPr.numId.val) for q in ps if isl(q)]
  ok([a for a,b in lv]==[0,1,1,0,0,0,0] and lv[0][1]==lv[1][1]==lv[3][1],'%svi_list ilvl/numId %s'%(pre,lv))
  s=ST[pre+'toc.pdf'];ok(s['status'].startswith('Done') and not s['ocrhint'],'%stoc no garbled warning'%pre)
  c=cer(f('toc'),os.path.join(T,'toc.docx'));ok(c<0.01,'%stoc CER %.4f'%(pre,c))
  ok(xml(f('toc')).count('w:leader="dot"')>=3,'%stoc dot leaders'%pre)
  s=ST[pre+'no_tounicode.pdf'];ok(s['status'].endswith('(OCR).') and s['href']=='/image-to-text/?evtest=1','%sOCR warning period + link %s'%(pre,s['href']))
  s=ST[pre+'toc_bad.pdf'];ok('garbled' in s['status'] and s['ocrhint'],'%stoc_bad garbled warning + OCR link'%pre)
  for n in ['vi_lo','vi_dejavu','Hợp đồng #1 (bản cuối) & thử','book','rot_a3','two_col','type3','vi_list','toc','numbers','cjk','pua_math','mixed','zh_latin']:ok(ST[pre+n+'.pdf']['status'].startswith('Done') and not ST[pre+n+'.pdf']['ocrhint'],'%s%s no warning'%(pre,n))
  for n in ['half_bad','few_bad','ten_bad','junkmap','caesar']:s=ST[pre+n+'.pdf'];ok('garbled' in s['status'] and s['ocrhint'],'%s%s garbled warning + OCR link'%(pre,n))
  c=cer(f('numbers'),os.path.join(T,'numbers.docx'));ok(c<0.01,'%snumbers CER %.4f'%(pre,c))
  c=cer(f('cjk'),os.path.join(T,'cjk.docx'));t=text(f('cjk'));ok(c<0.01 and 'しま す' not in t and '转换 质量' not in t and 'Word に' not in t,'%scjk CER %.4f no stray spaces'%(pre,c))
  import docx as _dx
  for n,k in [('cjk',2),('mixed',4),('zh_latin',4),('pua_math',None)]:
    c=cer(f(n),os.path.join(T,n+'.docx'));ok(c<0.005,'%s%s CER %.4f'%(pre,n,c))
    if k:np_=len([p for p in _dx.Document(f(n)).paragraphs if p.text.strip()]);ok(np_==k,'%s%s %d paragraphs (want %d)'%(pre,n,np_,k))
  t=text(f('mixed'));ok('Việt. こんにちは' in t and 'Hello, this is' in t,'%smixed spaces kept'%pre)
  t=text(f('zh_latin'));ok('使用 Microsoft' in t and '支持 JSON 格式' in t,'%szh_latin spaces kept'%pre)
  ok('通常 の' not in text(f('cjk')),'%scjk no space at line join'%pre)
  s=ST[pre+'vi_lo.pdf'];ok(s['status'].startswith('Done'),'%svi_lo status Done'%pre)
  ok(norm(text(f('type3')))!='','%stype3 text'%pre)
print('p2w fixtures:',len(fails),'failure(s)');sys.exit(1 if fails else 0)
